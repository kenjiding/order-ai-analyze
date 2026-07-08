# -*- coding: utf-8 -*-

import sys
import time
from datetime import datetime

import pandas as pd
import requests
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter


if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")


TOKEN = ""

OC_CENTER_CODE = "S210206"
COUNTRY = "AUS"
TIMEZONE = "+10"

PAGE_SIZE = 50000
WAYBILL_BATCH_SIZE = 1000

inventory_url = "https://oc.imile.com/saastms/v1/report/stock/daWaybillInventory"
waybill_url = "https://oc.imile.com/saastms/v1/biz/waybillInfo/searchWaybill"
organization_url = "https://oc.imile.com/lm/express/base/v1/biz/organization/getOrganizationByCode"

base_headers = {
    "authorization": f"Bearer {TOKEN}",
    "accept": "application/json, text/plain, */*",
    "content-type": "application/json",
    "front-sec": '{"orgId":10,"entId":10,"userCode":"2103269901","moduleId":10007,"client":"pc"}',
    "hubcode": OC_CENTER_CODE,
    "lang": "zh_CN",
    "origin": "https://oc.imile.com",
    "referer": "https://oc.imile.com/",
    "sys-name": "DS",
    "sys-portal": "OC",
    "timezone": TIMEZONE,
}

inventory_headers = {**base_headers, "resource-code": "DAInventory"}
waybill_headers = {**base_headers, "resource-code": "CentralWaybillQuery"}
organization_headers = {**base_headers, "resource-code": "OrganizationalManagement"}

request_count = {
    "仓外库存报表": 0,
    "中心运单查询": 0,
    "网点组织信息": 0,
}


def safe_json(response):
    try:
        return response.json()
    except Exception:
        return {"status": "failure", "message": response.text}


def batch_list(values, size):
    for i in range(0, len(values), size):
        yield values[i:i + size]


def find_result_object(data):
    return data.get("resultObject") or data.get("data") or data.get("result") or {}


def find_rows(result_object):
    if isinstance(result_object, list):
        return result_object

    if not isinstance(result_object, dict):
        return []

    for key in ["results", "records", "rows", "list", "dataList", "pageData", "items"]:
        value = result_object.get(key)
        if isinstance(value, list):
            return value

    for value in result_object.values():
        if isinstance(value, dict):
            rows = find_rows(value)
            if rows:
                return rows

    return []


def find_total_page(result_object, row_count, page_size):
    if not isinstance(result_object, dict):
        return 1

    pagination = result_object.get("pagination") or {}
    if isinstance(pagination, dict):
        total_page = pagination.get("totalPage")
        if total_page:
            return int(total_page)

    for key in ["totalPage", "pages"]:
        value = result_object.get(key)
        if value:
            return int(value)

    total = None
    for key in ["total", "totalCount", "count", "recordCount"]:
        value = result_object.get(key)
        if value:
            total = int(value)
            break

    if total:
        return max(1, (total + page_size - 1) // page_size)

    return 1 if row_count < page_size else 999999


def contains_dc(station_types):
    if station_types is None:
        return False

    if isinstance(station_types, str):
        return "DC" in station_types.upper()

    if isinstance(station_types, list):
        for item in station_types:
            if isinstance(item, str) and "DC" in item.upper():
                return True
            if isinstance(item, dict):
                text = " ".join(str(v) for v in item.values())
                if "DC" in text.upper():
                    return True

    return False


def find_route_value(row):
    for col in [
        "consigneeRouteCode",
        "routeCode",
        "dispatchRouteCode",
        "deliveryRouteCode",
        "route_code",
    ]:
        if col in row.index and pd.notna(row[col]) and str(row[col]).strip() != "":
            return row[col]
    return ""


def get_inventory_col(df, col_name):
    if col_name in df.columns:
        return df[col_name]
    return ""


def build_not_returned_range(value):
    if pd.isna(value):
        return ""

    value = int(value)

    if value < 0:
        return str(value)
    if value == 0:
        return "0D"
    if value == 1:
        return "1D"
    if value == 2:
        return "2D"
    if value == 3:
        return "3D"
    if value == 4:
        return "4D"
    if value == 5:
        return "5D"
    if value == 6:
        return "6D"
    if value == 7:
        return "7D"

    return ">7D"


def exclude_test_region(df):
    return df[~df["大区"].fillna("").astype(str).str.contains("test", case=False, na=False)].copy()


def build_station_summary(df):
    range_order = ["0D", "1D", "2D", "3D", "4D", "5D", "6D", "7D", ">7D"]
    output_columns = ["大区", "网点"] + range_order + ["Grand Total"]

    source = df.copy()
    source = source[source["最新扫描类型"].fillna("").astype(str).str.strip() == "派件"]
    source = exclude_test_region(source)
    source = source[source["未归班天范围"].fillna("").astype(str).isin(range_order)]

    if source.empty:
        return pd.DataFrame(columns=output_columns)

    pivot = pd.pivot_table(
        source,
        index=["大区", "stockStationName"],
        columns="未归班天范围",
        values="waybillNo",
        aggfunc="count",
        fill_value=0,
        observed=False,
    )

    for col in range_order:
        if col not in pivot.columns:
            pivot[col] = 0

    pivot = pivot[range_order]
    pivot["Grand Total"] = pivot.sum(axis=1)
    pivot = pivot.reset_index()

    rows = []

    for region in sorted(pivot["大区"].fillna("").astype(str).unique()):
        region_df = pivot[pivot["大区"].fillna("").astype(str) == region].copy()

        subtotal_row = {"大区": region or "未识别大区", "网点": ""}
        for col in range_order + ["Grand Total"]:
            subtotal_row[col] = int(region_df[col].sum())
        rows.append(subtotal_row)

        for _, row in region_df.sort_values("stockStationName").iterrows():
            detail_row = {"大区": "", "网点": row["stockStationName"]}
            for col in range_order + ["Grand Total"]:
                detail_row[col] = int(row[col])
            rows.append(detail_row)

    grand_total_row = {"大区": "Grand Total", "网点": ""}
    for col in range_order + ["Grand Total"]:
        grand_total_row[col] = int(pivot[col].sum())
    rows.append(grand_total_row)

    return pd.DataFrame(rows, columns=output_columns)


def build_driver_summary(df):
    range_order = ["2D", "3D", "4D", "5D", "6D", "7D", ">7D"]
    output_columns = ["大区", "网点", "司机", "供应商", "路由"] + range_order + ["Grand Total"]

    source = df.copy()
    source = source[source["最新扫描类型"].fillna("").astype(str).str.strip() == "派件"]
    source = exclude_test_region(source)
    source = source[source["未归班天范围"].fillna("").astype(str).isin(range_order)]

    if source.empty:
        return pd.DataFrame(columns=output_columns)

    source["driverName"] = source["driverName"].fillna("").astype(str).str.strip()
    source["driverVendorName"] = source["driverVendorName"].fillna("").astype(str).str.strip()
    source["路由码"] = source["路由码"].fillna("").astype(str).str.strip()

    driver_total = (
        source.groupby(["大区", "stockStationName", "driverName", "driverVendorName"], dropna=False)["waybillNo"]
        .count()
        .reset_index(name="driver_total")
    )

    qualified_driver = driver_total[driver_total["driver_total"] > 10].copy()

    if qualified_driver.empty:
        return pd.DataFrame(columns=output_columns)

    source = source.merge(
        qualified_driver[["大区", "stockStationName", "driverName", "driverVendorName"]],
        on=["大区", "stockStationName", "driverName", "driverVendorName"],
        how="inner",
    )

    pivot = pd.pivot_table(
        source,
        index=["大区", "stockStationName", "driverName", "driverVendorName", "路由码"],
        columns="未归班天范围",
        values="waybillNo",
        aggfunc="count",
        fill_value=0,
        observed=False,
    )

    for col in range_order:
        if col not in pivot.columns:
            pivot[col] = 0

    pivot = pivot[range_order]
    pivot["Grand Total"] = pivot.sum(axis=1)
    pivot = pivot.reset_index()

    pivot = pivot.sort_values(
        ["大区", "stockStationName", "driverName", "driverVendorName", "Grand Total"],
        ascending=[True, True, True, True, False],
    )

    pivot = pivot.rename(
        columns={
            "stockStationName": "网点",
            "driverName": "司机",
            "driverVendorName": "供应商",
            "路由码": "路由",
        }
    )

    return pivot[output_columns]


def merge_same_cells_by_key(sheet, target_col, key_cols, start_row=3):
    max_row = sheet.max_row
    if max_row < start_row:
        return

    cached_rows = {}
    for row_idx in range(start_row, max_row + 1):
        cached_rows[row_idx] = {
            col_idx: sheet.cell(row=row_idx, column=col_idx).value
            for col_idx in key_cols
        }

    group_start = start_row
    previous_key = tuple(cached_rows[start_row][col_idx] for col_idx in key_cols)

    for row_idx in range(start_row + 1, max_row + 2):
        if row_idx <= max_row:
            current_key = tuple(cached_rows[row_idx][col_idx] for col_idx in key_cols)
        else:
            current_key = None

        if current_key != previous_key:
            group_end = row_idx - 1

            if group_end > group_start:
                sheet.merge_cells(
                    start_row=group_start,
                    start_column=target_col,
                    end_row=group_end,
                    end_column=target_col,
                )

                cell = sheet.cell(row=group_start, column=target_col)
                cell.alignment = Alignment(horizontal="center", vertical="center")

            group_start = row_idx
            previous_key = current_key


def merge_driver_summary_cells(sheet):
    # 大区相同合并
    merge_same_cells_by_key(sheet, target_col=1, key_cols=[1], start_row=3)

    # 同一大区 + 网点相同合并
    merge_same_cells_by_key(sheet, target_col=2, key_cols=[1, 2], start_row=3)

    # 同一大区 + 网点 + 司机相同合并
    merge_same_cells_by_key(sheet, target_col=3, key_cols=[1, 2, 3], start_row=3)


def style_excel(file_path):
    workbook = load_workbook(file_path)

    header_fill = PatternFill("solid", fgColor="5B7FA6")
    title_fill = PatternFill("solid", fgColor="1F4E78")
    region_fill = PatternFill("solid", fgColor="D9EAF7")
    total_fill = PatternFill("solid", fgColor="2F3A4A")

    white_font = Font(color="FFFFFF", bold=True)
    title_font = Font(color="FFFFFF", bold=True, size=14)
    bold_font = Font(color="000000", bold=True)

    border = Border(
        left=Side(style="thin", color="B7B7B7"),
        right=Side(style="thin", color="B7B7B7"),
        top=Side(style="thin", color="B7B7B7"),
        bottom=Side(style="thin", color="B7B7B7"),
    )

    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]

        if sheet_name in ["网点汇总", "司机汇总"]:
            sheet.freeze_panes = "A3"

            max_col = sheet.max_column
            sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)

            title_cell = sheet.cell(row=1, column=1)
            title_cell.fill = title_fill
            title_cell.font = title_font
            title_cell.alignment = Alignment(horizontal="center", vertical="center")
            title_cell.border = border
            sheet.row_dimensions[1].height = 30

            header_row = 2
        else:
            sheet.freeze_panes = "A2"
            header_row = 1

        sheet.auto_filter.ref = sheet.dimensions

        for cell in sheet[header_row]:
            cell.fill = header_fill
            cell.font = white_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

        for row in sheet.iter_rows():
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="center")

        if sheet_name == "网点汇总":
            for row_idx in range(3, sheet.max_row + 1):
                region_value = sheet.cell(row=row_idx, column=1).value
                is_grand_total = region_value == "Grand Total"
                is_region_total = region_value not in [None, "", "Grand Total"]

                for col_idx in range(1, sheet.max_column + 1):
                    cell = sheet.cell(row=row_idx, column=col_idx)

                    if col_idx >= 3:
                        cell.alignment = Alignment(horizontal="right", vertical="center")

                    if is_grand_total:
                        cell.fill = total_fill
                        cell.font = white_font
                    elif is_region_total:
                        cell.fill = region_fill
                        cell.font = bold_font

        if sheet_name == "司机汇总":
            merge_driver_summary_cells(sheet)

            for row_idx in range(3, sheet.max_row + 1):
                for col_idx in range(6, sheet.max_column + 1):
                    sheet.cell(row=row_idx, column=col_idx).alignment = Alignment(horizontal="right", vertical="center")

        for col_idx, column_cells in enumerate(sheet.columns, start=1):
            max_length = 0
            for cell in column_cells:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))
            sheet.column_dimensions[get_column_letter(col_idx)].width = min(max(max_length + 2, 10), 30)

    workbook.save(file_path)


# 1. 查询仓外库存，全量不传 start/end
all_rows = []

payload = {
    "currentPage": 1,
    "showCount": PAGE_SIZE,
    "orderTypes": [],
    "stockReportType": "allTotal",
    "scanType": "",
    "scanTypeList": [],
    "stockStationCodes": [],
    "ocCenterCode": OC_CENTER_CODE,
    "country": COUNTRY,
    "fromOc": True,
}

r = requests.post(inventory_url, headers=inventory_headers, json=payload, timeout=180)
request_count["仓外库存报表"] += 1
data = safe_json(r)

result_object = find_result_object(data)
first_rows = find_rows(result_object)

if not result_object and not first_rows:
    raise Exception(f"仓外库存报表接口返回异常：{data}")

total_page = find_total_page(result_object, len(first_rows), PAGE_SIZE)
print(f"仓外库存报表总页数：{total_page}")

all_rows.extend(first_rows)
print(f"第 1 页返回：{len(first_rows)} 条")

for page in range(2, total_page + 1):
    payload["currentPage"] = page

    r = requests.post(inventory_url, headers=inventory_headers, json=payload, timeout=180)
    request_count["仓外库存报表"] += 1
    data = safe_json(r)

    result_object = find_result_object(data)
    rows = find_rows(result_object)

    if not rows:
        print(f"第 {page} 页没有返回明细，停止分页")
        break

    all_rows.extend(rows)
    print(f"第 {page} 页返回：{len(rows)} 条，目前共 {len(all_rows)} 条")

    if len(rows) < PAGE_SIZE:
        break

    time.sleep(0.2)

df_inventory = pd.DataFrame(all_rows)

if df_inventory.empty:
    raise Exception("没有抓到任何数据，请检查 TOKEN、网点编码。")

print(f"仓外库存记录数：{len(df_inventory)}")

if "stockStationCode" not in df_inventory.columns:
    raise Exception("接口结果里没有 stockStationCode 字段，无法准确查询大区。")


# 2. 整理基础字段
output_df = pd.DataFrame()
output_df["stockStationCode"] = get_inventory_col(df_inventory, "stockStationCode")
output_df["stockStationName"] = get_inventory_col(df_inventory, "stockStationName")
output_df["ofdDate"] = get_inventory_col(df_inventory, "ofdDate")
output_df["waybillNo"] = get_inventory_col(df_inventory, "waybillNo")
output_df["driverName"] = get_inventory_col(df_inventory, "driverName")
output_df["driverVendorName"] = get_inventory_col(df_inventory, "driverVendorName")
output_df["clientName"] = get_inventory_col(df_inventory, "clientName")
output_df["deliveryStationName"] = get_inventory_col(df_inventory, "deliveryStationName")
output_df["deliveryCount"] = get_inventory_col(df_inventory, "deliveryCount")
output_df["最新扫描类型"] = get_inventory_col(df_inventory, "lastScanTypeDesc")

if output_df["waybillNo"].fillna("").astype(str).str.strip().eq("").all():
    raise Exception("接口结果里没有有效 waybillNo，无法查询路由码。")


# 3. 用 waybillNo 查询中心运单，补路由码
waybill_list = (
    output_df["waybillNo"]
    .dropna()
    .astype(str)
    .str.strip()
    .replace("", pd.NA)
    .dropna()
    .drop_duplicates()
    .tolist()
)

print(f"准备查询路由码运单数：{len(waybill_list)}")

route_map = {}
failed_waybill_batches = []

max_retry = 3
start_time = "2020-01-01 00:00:00"
end_time = datetime.now().strftime("%Y-%m-%d 23:59:59")

waybill_batches = list(batch_list(waybill_list, WAYBILL_BATCH_SIZE))
total_waybill_batches = len(waybill_batches)

for batch_index, batch in enumerate(waybill_batches, start=1):
    waybill_payload = {
        "currentPage": 1,
        "showCount": WAYBILL_BATCH_SIZE,
        "waybillNos": batch,
        "referenceNos": [],
        "labelOrderCodes": [],
        "mainWaybillNos": [],
        "platformNos": [],
        "dateSearchType": "waybill",
        "startTime": start_time,
        "endTime": end_time,
        "actualReceiveStation": "",
    }

    success = False

    for retry in range(1, max_retry + 1):
        try:
            r = requests.post(waybill_url, headers=waybill_headers, json=waybill_payload, timeout=60)
            request_count["中心运单查询"] += 1
            data = safe_json(r)

            result_object = find_result_object(data)
            rows = find_rows(result_object)

            if not rows:
                print(f"中心运单第 {batch_index}/{total_waybill_batches} 批无数据，第 {retry} 次")
                time.sleep(1)
                continue

            df_waybill_batch = pd.DataFrame(rows)

            if "waybillNo" in df_waybill_batch.columns:
                for _, row in df_waybill_batch.iterrows():
                    waybill_no = str(row.get("waybillNo", "")).strip()
                    route_code = find_route_value(row)
                    if waybill_no:
                        route_map[waybill_no] = route_code

            success = True
            break

        except Exception as e:
            print(f"中心运单第 {batch_index}/{total_waybill_batches} 批异常，第 {retry} 次：{e}")
            time.sleep(1)

    if not success:
        failed_waybill_batches.append(batch)

    print(f"路由码进度：{batch_index}/{total_waybill_batches} 批，已匹配 {len(route_map)} 条")
    time.sleep(0.2)

output_df["路由码"] = output_df["waybillNo"].astype(str).str.strip().map(route_map).fillna("")


# 4. 用 stockStationCode 查询网点组织信息，补大区
station_info_df = output_df[["stockStationCode", "stockStationName"]].copy()
station_info_df["stockStationCode"] = station_info_df["stockStationCode"].fillna("").astype(str).str.strip()
station_info_df["stockStationName"] = station_info_df["stockStationName"].fillna("").astype(str).str.strip()
station_info_df = station_info_df.drop_duplicates()

region_map_by_code = {}
region_map_by_name = {}

for idx, row in station_info_df.reset_index(drop=True).iterrows():
    stock_station_code = str(row["stockStationCode"]).strip()
    stock_station_name = str(row["stockStationName"]).strip()

    if not stock_station_code:
        continue

    try:
        r = requests.get(
            organization_url,
            headers=organization_headers,
            params={"code": stock_station_code},
            timeout=30,
        )
        request_count["网点组织信息"] += 1
        data = safe_json(r)

        org_info = data.get("resultObject") or data.get("data") or {}

        organization_name = (
            org_info.get("organizationName")
            or org_info.get("stationName")
            or stock_station_name
        )

        parent_station_name = (
            org_info.get("parent_station_name")
            or org_info.get("parentStationName")
            or ""
        )

        station_types = (
            org_info.get("station_types")
            or org_info.get("stationTypes")
            or []
        )

        if contains_dc(station_types):
            region_name = stock_station_name
        else:
            region_name = parent_station_name

        print(
            f"大区匹配：{stock_station_name} / {stock_station_code} "
            f"-> parentStationName={parent_station_name}, stationTypes={station_types}, 大区={region_name}"
        )

        if region_name:
            region_map_by_code[stock_station_code] = region_name
            region_map_by_name[stock_station_name] = region_name
            region_map_by_name[organization_name] = region_name

    except Exception as e:
        print(
            f"网点组织信息查询异常，"
            f"stockStationName={stock_station_name}, stockStationCode={stock_station_code}：{e}"
        )

    if (idx + 1) % 20 == 0 or (idx + 1) == len(station_info_df):
        print(f"网点组织信息进度：{idx + 1}/{len(station_info_df)}")

    time.sleep(0.1)

output_df["大区"] = output_df["stockStationCode"].fillna("").astype(str).str.strip().map(region_map_by_code).fillna("")

empty_region_mask = output_df["大区"].fillna("").astype(str).str.strip() == ""

output_df.loc[empty_region_mask, "大区"] = (
    output_df.loc[empty_region_mask, "stockStationName"]
    .fillna("")
    .astype(str)
    .str.strip()
    .map(region_map_by_name)
    .fillna("")
)


# 5. 添加未归班天数、未归班天范围
today_date = pd.to_datetime(datetime.now().strftime("%Y-%m-%d"))
ofd_date = pd.to_datetime(output_df["ofdDate"], errors="coerce")

output_df["未归班天数"] = (today_date - ofd_date.dt.normalize()).dt.days
output_df["未归班天范围"] = output_df["未归班天数"].apply(build_not_returned_range)


# 6. 整理字段顺序
output_df = output_df[
    [
        "大区",
        "stockStationName",
        "ofdDate",
        "未归班天数",
        "未归班天范围",
        "waybillNo",
        "路由码",
        "最新扫描类型",
        "driverName",
        "driverVendorName",
        "clientName",
        "deliveryStationName",
        "deliveryCount",
    ]
]


# 7. 生成汇总
station_summary_table = build_station_summary(output_df)
driver_summary_table = build_driver_summary(output_df)


# 8. 导出 Excel
file_name = datetime.now().strftime("%Y%m%d%H%M")
today_text = datetime.now().strftime("%Y-%m-%d")
output_file = rf"C:\Users\iMileUSER\Downloads\{file_name}_DAInventory_Selected_Rawdata.xlsx"

with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
    output_df.to_excel(writer, sheet_name="Rawdata", index=False)

    station_summary_table.to_excel(
        writer,
        sheet_name="网点汇总",
        index=False,
        startrow=1,
    )

    driver_summary_table.to_excel(
        writer,
        sheet_name="司机汇总",
        index=False,
        startrow=1,
    )

    workbook = writer.book
    workbook["网点汇总"].cell(row=1, column=1).value = f"{today_text} 仓外库存"
    workbook["司机汇总"].cell(row=1, column=1).value = f"{today_text} 超2天&超10件派件未归班"

style_excel(output_file)

if failed_waybill_batches:
    failed_file = rf"C:\Users\iMileUSER\Downloads\{file_name}_failed_waybill_batches.txt"
    with open(failed_file, "w", encoding="utf-8") as f:
        for batch in failed_waybill_batches:
            f.write(",".join(batch) + "\n")
    print(f"失败批次已保存：{failed_file}")

print(f"导出完成：{output_file}")
print(f"导出记录数：{len(output_df)}")
print(f"路由码匹配数：{output_df['路由码'].fillna('').astype(str).str.strip().ne('').sum()}")
print(f"大区匹配数：{output_df['大区'].fillna('').astype(str).str.strip().ne('').sum()}")
print(f"网点汇总行数：{len(station_summary_table)}")
print(f"司机汇总行数：{len(driver_summary_table)}")
print(f"请求次数：{request_count}")