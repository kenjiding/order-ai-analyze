#!/usr/bin/env python3
"""Generate internal AI pioneer showcase Word document from ai分析 images."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE_DIR = Path(__file__).resolve().parent
OUTPUT = BASE_DIR / "AI先行者实践案例-仓库数据分析与异常研判.docx"

IMAGES = {
    "warehouse": BASE_DIR / "仓库整体效率.png",
    "warehouse_action": BASE_DIR / "根据仓内整体效率-进行调整.jpg",
    "scanner": BASE_DIR / "分拣员个人效率.png",
    "scanner_action": BASE_DIR / "根据分拣员个人效率-进行单独培训.jpg",
    "order_analyze": BASE_DIR / "AI异常订单分析.png",
    "order_push": BASE_DIR / "ai异常订单分析推送.png",
}


def set_doc_font(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "PingFang SC"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")


def add_title(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)


def add_para(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.35


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.35


def add_image(doc: Document, path: Path, caption: str, width_cm: float = 16.0) -> None:
    if not path.exists():
        add_para(doc, f"[图片缺失: {path.name}]", bold=True)
        return
    doc.add_picture(str(path), width=Cm(width_cm))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(9)
    cap.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    doc.add_paragraph()


def build_document() -> Document:
    doc = Document()
    set_doc_font(doc)

    # Cover
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("AI 先行者实践案例\n仓库数据分析与异常研判")
    tr.bold = True
    tr.font.size = Pt(22)
    tr.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("—— 用真实业务场景推广 AI 落地 ——")
    sr.font.size = Pt(12)
    sr.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    doc.add_paragraph()
    add_para(
        doc,
        "本文档整理自团队自研「扫描效率分析」与「AI 订单异常管理」系统的真实截图，"
        "展示 AI 如何把考勤、扫描、轨迹等结构化数据转化为可执行的现场管理建议。"
        "适合在公司内部群分享，作为 AI 先行者的模仿示例，帮助同事理解：AI 不是概念，而是能直接辅助排班、培训与异常处置的工具。",
    )

    add_title(doc, "一、我们在做什么", 1)
    add_para(
        doc,
        "仓库运营每天产生大量数据：开仓货量与工时、分拣员扫描记录、异常订单轨迹。"
        "传统做法是人工翻表、凭经验判断；我们的做法是：系统先算好指标，再让 AI 读数据、写结论，"
        "管理者只需核对数字、落地动作。",
    )
    add_bullets(
        doc,
        [
            "开仓效率 / 总体效率：看全仓产能与成本，发现异常日、波动原因。",
            "分拣员个人效率：看个人与团队差距，精准安排培训与排班。",
            "AI 异常订单分析：批量读轨迹，自动判责、分类，减少人工逐单排查。",
        ],
    )

    add_title(doc, "二、开仓效率 / 总体效率 —— 从数据到排班调整", 1)
    add_title(doc, "2.1 AI 分析逻辑", 2)
    add_para(doc, "系统向 AI 输入以下结构化数据（口径与页面一致）：", bold=True)
    add_bullets(
        doc,
        [
            "开仓效率（票/小时）：开仓日货量 ÷ 开仓日工时。",
            "总体效率（票/小时）：总货量 ÷ 全部工时（含非开仓时段）。",
            "按日明细：日期、货量、工时、是否开仓日、当日效率；最多 20 条。",
            "辅助指标：开仓天数、统计总天数、菜鸟占比、每票成本等（页面展示）。",
        ],
    )
    add_para(
        doc,
        "AI 输出两类内容：findings（3～5 条数据亮点或问题）与 suggestions（2～4 条具体改进建议）。"
        "语气像现场主管写给管理层的备注：直接扣数字，禁止空泛套话。",
    )

    add_title(doc, "2.2 系统截图示例", 2)
    add_image(
        doc,
        IMAGES["warehouse"],
        "图 1  开仓效率 / 总体效率页面：KPI、AI 数据分析建议、按日效率明细",
    )
    add_para(doc, "本例 AI 解读要点（摘自截图）：", bold=True)
    add_bullets(
        doc,
        [
            "开仓效率 148.10 与总体效率 146.91 接近，说明非开仓工时占比很小（约 7.75 小时），对总效率影响有限。",
            "6 月 13 日开仓效率仅 2.19 票/小时（57 票 / 26 工时），属严重异常，需排查人员、设备或流程。",
            "6 月 12 日峰值 179.57 票/小时，可复盘当日排班与货量结构，提炼可复制做法。",
            "6 月 5 日效率 111.66 低于均值，建议结合货型复杂度、人手配置进一步分析。",
        ],
    )

    add_title(doc, "2.3 落地动作：根据仓内整体效率调整现场", 2)
    add_para(
        doc,
        "AI 结论不是停在屏幕上。管理者根据 findings 调整排班、货量节奏与异常日响应——"
        "例如异常日加派人手、峰值日固化排班模板、低效日检查货型与设备。",
    )
    add_image(
        doc,
        IMAGES["warehouse_action"],
        "图 2  现场分拣：依据整体效率分析，优化开仓节奏、笼车分配与货量堆积管理",
        width_cm=14.5,
    )

    add_title(doc, "三、分拣员个人效率 —— 从排名到一对一培训", 1)
    add_title(doc, "3.1 AI 分析逻辑", 2)
    add_para(doc, "系统向 AI 输入分拣侧结构化数据：", bold=True)
    add_bullets(
        doc,
        [
            "一级分拣效率（票/小时）：总扫描量 ÷ 有效工时。",
            "扫描员明细：姓名、扫描量、有效工时、原始效率、修正效率（adjustedEfficiency）。",
            "修正效率说明：样本量或日均工时不足时，向团队均值回归，避免短班、小样本误判。",
            "汇总：扫描员人数、统计天数、总扫描量、总有效工时。",
        ],
    )
    add_para(
        doc,
        "AI 同样输出 findings + suggestions，重点识别人效标杆与拖后腿人员，并给出培训、排班、激励建议。",
    )

    add_title(doc, "3.2 系统截图示例", 2)
    add_image(
        doc,
        IMAGES["scanner"],
        "图 3  分拣员个人效率页面：团队 KPI、AI 建议、扫描员效率汇总与总效率对比",
    )
    add_para(doc, "本例 AI 解读要点（摘自截图）：", bold=True)
    add_bullets(
        doc,
        [
            "团队一级分拣效率 545 票/小时，低于 600 的行业参考线，仍有提升空间。",
            "头部 5 名与尾部 5 名相差近 200 票/小时，人效分化明显。",
            "ADL-DC-shigong03、Tong、Bartel、Kuanhua 等原始效率低于 400，建议一对一辅导。",
            "ADL-DC-Miko、ADL-YU WUN 等超过 680，可作为标杆组织经验分享。",
            "建议：低效人员非高峰时段排班、每日效率榜、标杆带教扫描节奏与异常件处理。",
        ],
    )

    add_title(doc, "3.3 落地动作：根据个人效率安排单独培训", 2)
    add_para(
        doc,
        "个人效率分析的价值在于「点名到人」：不是笼统喊「大家加油」，"
        "而是对尾部扫描员做跟班、录屏复盘、标杆结对；对头部员工沉淀 SOP 供新人学习。",
    )
    add_image(
        doc,
        IMAGES["scanner_action"],
        "图 4  现场带教：依据个人效率数据，对低效分拣员进行一对一操作辅导",
        width_cm=14.5,
    )

    add_title(doc, "四、AI 异常订单分析 —— 从逐单排查到批量判责", 1)
    add_title(doc, "4.1 AI 分析逻辑", 2)
    add_para(doc, "与效率分析不同，异常订单分析走的是「轨迹研判」流水线：", bold=True)
    add_bullets(
        doc,
        [
            "输入：订单号列表 → 系统拉取 iMile 最新轨迹与阻断日志（freshForAi）。",
            "预处理：服务端计算时序分析（责任仓、库内/库外起算、超时阈值、真实操作员识别等），模型只读预计算字段，不自行猜规则。",
            "AI 输出（每单）：丢失责任类型（库内丢失 / 库外丢失 / 需人工判断等）、子类型、判定依据、置信度。",
            "护栏：服务端硬规则与模型结论交叉校验；不确定的订单标记需人工复核。",
        ],
    )
    add_para(
        doc,
        "单次最多 500 单；支持 DeepSeek / OpenAI 切换；大批量可走异步任务队列，结果可导出 Excel。",
    )

    add_title(doc, "4.2 系统截图：批量分析与结果表", 2)
    add_image(
        doc,
        IMAGES["order_analyze"],
        "图 5  AI 订单异常管理：粘贴单号 → 批量分析 → 异常类型、责任方、置信度、判定依据",
    )

    add_title(doc, "4.3 系统截图：飞书推送日报", 2)
    add_image(
        doc,
        IMAGES["order_push"],
        "图 6  AI 订单异常分析报告（飞书卡片）：汇总 + 类型分布 + 责任方 Top8 + 高频子分类",
    )
    add_para(doc, "本例推送解读要点（摘自截图，2026-06-19）：", bold=True)
    add_bullets(
        doc,
        [
            "共分析 103 单异常：库内丢失 55 单（53.4%），库外丢失 27 单（26.2%）。",
            "责任站点 Top：Sydney Center 33 单、Melbourne Center 18 单。",
            "高频子类：扫描员丢失 55 单、快递员丢失 27 单。",
            "AI 建议：重点盯 Sydney / Melbourne 扫描规范与派送环节防丢措施。",
        ],
    )

    add_title(doc, "五、三条业务线如何串起来", 1)
    add_para(
        doc,
        "效率分析与异常分析解决的是同一仓库的不同问题，但可以形成管理闭环：",
    )
    add_bullets(
        doc,
        [
            "整体效率看「仓有没有开够、哪天出了问题」→ 调整排班与开仓节奏（图 1→2）。",
            "个人效率看「谁需要帮、谁可以带」→ 精准培训与激励（图 3→4）。",
            "异常订单看「丢在哪、谁的责任」→ 减少人工逐单看轨迹，把精力放到 Sydney/MEL 等高发站点（图 5→6）。",
        ],
    )
    add_para(
        doc,
        "异常分析里「扫描员丢失」高发，与个人效率里尾部扫描员低效，可以交叉验证："
        "既做操作培训，也加强扫描规范与防丢流程。",
    )

    add_title(doc, "六、给全体同事的一句话", 1)
    add_para(
        doc,
        "AI 先行不是要写多复杂的 Prompt，而是：先把业务数据算对、口径统一，"
        "再让 AI 在固定结构里「读数 + 写建议 + 批量判责」。"
        "欢迎各区域、各岗位参照以上三个场景，从「一个小痛点 + 一份真实数据」开始尝试。"
        "有想法或需要接入数据的，可以直接找本案例负责人交流。",
    )

    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("—— 内部资料，请勿对外传播 ——")
    fr.font.size = Pt(9)
    fr.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    return doc


def main() -> None:
    doc = build_document()
    doc.save(str(OUTPUT))
    print(f"Generated: {OUTPUT}")


if __name__ == "__main__":
    main()
